"""Xuất phụ lục cho thư mục nộp 05_Phu_luc (chạy tự động trong run_all.py).

  05_Phu_luc/Nhat_ky_ETL.docx      <- data/processed/etl_log.json + 05_Phu_luc/ghi_chu_thu_cong.md
  05_Phu_luc/Tu_dien_du_lieu.xlsx  <- 05_Phu_luc/Tu_dien_du_lieu.md (mỗi mục "## ..." thành 1 sheet)

Ghi chú thủ công (đối chiếu, giải thích ngoại lệ) viết trong ghi_chu_thu_cong.md – file này
KHÔNG bị ghi đè khi chạy lại run_all.py.
Cần: pip install python-docx openpyxl
"""
from __future__ import annotations

import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
DOC_DIR = ROOT / "05_Phu_luc"
FONT = "Arial"


# ------------------------------------------------------------------ markdown helpers
def _md_tables_and_bullets(lines: list[str]) -> list[list[str]]:
    """Trả về các dòng dạng list ô: bảng markdown -> nhiều ô; gạch đầu dòng/đoạn -> 1 ô."""
    rows = []
    for ln in lines:
        s = ln.strip()
        if not s or re.fullmatch(r"\|?[\s:\-|]+\|?", s):
            continue
        if s.startswith("|"):
            rows.append([c.strip() for c in s.strip("|").split("|")])
        else:
            rows.append([re.sub(r"^[-*]\s+", "", s)])
    return rows


def _plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*|`(.+?)`", lambda m: m.group(1) or m.group(2), text)


def _sections(md: str) -> list[tuple[str, list[str]]]:
    out, title, buf = [], "Gioi_thieu", []
    for ln in md.splitlines():
        if ln.startswith("## "):
            out.append((title, buf))
            title, buf = ln[3:].strip(), []
        elif ln.startswith("# "):
            continue
        else:
            buf.append(ln)
    out.append((title, buf))
    return [(t, b) for t, b in out if any(x.strip() for x in b)]


# ------------------------------------------------------------------ Từ điển -> xlsx
def tu_dien_xlsx(md_path: Path = DOC_DIR / "Tu_dien_du_lieu.md", out: Path = DOC_DIR / "Tu_dien_du_lieu.xlsx") -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    md = md_path.read_text(encoding="utf-8")
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BFBFBF")
    head_fill = PatternFill("solid", fgColor="1F3864")
    todo_fill = PatternFill("solid", fgColor="FFF2CC")
    for title, lines in _sections(md):
        name = re.sub(r"\s*\(.*?\)", "", re.sub(r"^\d+\.\s*", "", title))[:31].replace("/", "-")
        ws = wb.create_sheet(name or "Sheet")
        rows = _md_tables_and_bullets(lines)
        is_table = any(len(r) > 1 for r in rows)
        width = max(len(r) for r in rows) if rows else 1
        for i, r in enumerate(rows, 1):
            for j, v in enumerate(r, 1):
                c = ws.cell(i, j, _plain(v))
                c.font = Font(name=FONT, size=10, bold=is_table and i == 1, color="FFFFFF" if is_table and i == 1 else "000000")
                c.alignment = Alignment(wrap_text=True, vertical="top")
                c.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if is_table and i == 1:
                    c.fill = head_fill
                elif "____" in str(v) or "__/" in str(v):
                    c.fill = todo_fill   # ô vàng = còn phải điền
        for j in range(1, width + 1):
            col = ws.cell(1, j).column_letter
            longest = max((len(str(ws.cell(i, j).value or "")) for i in range(1, ws.max_row + 1)), default=10)
            ws.column_dimensions[col].width = min(max(12, longest * 0.9), 70)
        ws.freeze_panes = "A2" if is_table else None
    note = wb.create_sheet("Ghi_chu")
    note["A1"] = "Sinh tự động từ 05_Phu_luc/Tu_dien_du_lieu.md bằng src/export_phu_luc.py. Ô vàng = còn phải điền. Sửa file .md rồi chạy lại run_all.py."
    note["A1"].font = Font(name=FONT, size=10, italic=True)
    wb.save(out)
    return out


# ------------------------------------------------------------------ Nhật ký -> docx
def nhat_ky_docx(log_path: Path = ROOT / "data" / "processed" / "etl_log.json",
                 notes_path: Path = DOC_DIR / "ghi_chu_thu_cong.md",
                 out: Path = DOC_DIR / "Nhat_ky_ETL.docx") -> Path:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    log = json.loads(log_path.read_text(encoding="utf-8"))
    d = Document()
    st = d.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    for h in ("Heading 1", "Heading 2"):
        d.styles[h].font.name = FONT
        d.styles[h].font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    def shade(cell, hexcolor):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), hexcolor)
        tcPr.append(shd)

    def table(header, rows):
        t = d.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, h in zip(t.rows[0].cells, header):
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            shade(c, "1F3864")
        for row in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, row):
                c.text = str(v)
                if str(v).startswith("LOẠI"):
                    shade(c, "FCE4D6")
        return t

    def bullet(text):
        d.add_paragraph(text, style="List Bullet")

    d.add_heading("Nhật ký ETL – Bước O (dữ liệu)", 1)
    d.add_paragraph("Đề tài 12 – Tối ưu hóa danh mục cổ phiếu Việt Nam. Người phụ trách: TV2. "
                    "Sinh tự động bởi run_all.py từ data/processed/etl_log.json.")
    cal = log["calendar"]
    d.add_heading("1. Tổng quan", 2)
    table(["Mục", "Giá trị"], [
        ["Cửa sổ dữ liệu", f"{log['window'][0]} → {log['window'][1]}"],
        ["Lịch giao dịch chung (theo VN-Index)", f"{cal['trading_days']} phiên ({cal['first']} → {cal['last']})"],
        ["Vũ trụ ban đầu / giữ lại", f"{log['universe_size']} mã (VN30 kỳ 01/2020) / {log['final_shape']['tickers']} mã"],
        ["Quy tắc điền thiếu", f"Điền giá phiên trước tối đa {log['max_ffill']} phiên + cột cờ imputed"],
        ["Ô giá được điền (imputed)", str(log["final_imputed_cells"])],
        ["Ô NaN còn lại", f"{log['final_na_cells']} (thuộc các sự kiện ở mục 5)"],
        ["VN-Index", f"Nguồn: {log.get('vnindex_source')}; khoảng {log.get('vnindex_range')} điểm"
                     + (" – ĐỘ CHÍNH XÁC THẤP" if log.get("vnindex_low_precision") else "")],
    ])

    d.add_heading("2. Quy tắc làm sạch (4 bước)", 2)
    for s in [
        "Bước 1 – Chuẩn hóa: tên cột chữ thường; ngày về dạng ngày thuần (bỏ múi giờ +07:00); bỏ ngày trùng (giữ bản cuối); "
        "giá ≤ 0 → trống; trung vị giá < 1.000 → nhân 1.000 (nghìn đồng → VND).",
        "Bước 2 – Căn lịch: mọi mã đưa về cùng lịch phiên giao dịch của VN-Index.",
        f"Bước 3 – Xử lý thiếu: điền giá phiên trước tối đa {log['max_ffill']} phiên (chỉ trong khoảng mã còn giao dịch), "
        "gắn cờ imputed = 1; thiếu dài hơn để trống và ghi lại.",
        "Bước 4 – Kiểm tra cuối: loại mã thiếu lịch sử hoặc thiếu dài trước kỳ quyết định đầu; mã ngừng giao dịch sau đó "
        "được GIỮ (tránh thiên lệch sống sót) và ghi vào events.csv; quét biến động ngày > 7,5% (vượt biên độ ±7% HOSE) "
        "để phát hiện giá chưa điều chỉnh.",
    ]:
        bullet(s)

    d.add_heading("3. Số dòng trước/sau theo mã", 2)
    rows = []
    for t, s in log["per_ticker"].items():
        status = "LOẠI: " + log["excluded"][t] if t in log["excluded"] else "giữ"
        rows.append([t, s["rows_raw"], s["dup_dates"], s["nonpositive"], "có" if s["scaled_x1000"] else "không",
                     s["rows_after_calendar"], s.get("imputed_cells", 0), status])
    table(["Mã", "Dòng thô", "Trùng ngày", "Giá ≤0", "×1000", "Sau căn lịch", "Imputed", "Trạng thái"], rows)

    d.add_heading("4. Mã bị loại", 2)
    for t, r in log["excluded"].items() or {}.items():
        bullet(f"{t}: {r}")
    if not log["excluded"]:
        bullet("(không có)")

    d.add_heading("5. Sự kiện giữ lại trong dữ liệu (data/processed/events.csv)", 2)
    ev = log.get("events", [])
    if ev:
        table(["Mã", "Sự kiện", "Từ", "Đến", "Số phiên", "Giá cuối (VND)", "Xử lý"],
              [[e["ticker"], e["event"], e["start"], e["end"] or "—", e["sessions"],
                "" if e["last_close"] != e["last_close"] else f"{e['last_close']:,.0f}", e["note"]] for e in ev])
    else:
        bullet("(không có)")

    d.add_heading("6. Biến động ngày > 7,5% (cần giải thích)", 2)
    sus = log.get("suspect_unadjusted", {})
    for t, v in sus.items():
        bullet(f"{t}: " + "; ".join(f"{a} ({b:+.1%})" for a, b in v))
    if not sus:
        bullet("(không có) → không phát hiện dấu hiệu giá chưa điều chỉnh")

    d.add_heading("7. Lãi suất phi rủi ro (TPCP 10 năm)", 2)
    rf = log["rf"]
    table(["Chỉ tiêu", "Giá trị"], [
        ["Số dòng nguồn", rf["rows_source"]],
        ["Số phiên có rf / tổng", f"{rf['calendar_days'] - rf['na_after_fill']} / {rf['calendar_days']}"],
        ["rf thấp nhất – cao nhất (%/năm)", f"{rf['rf_annual_min'] * 100:.2f}% – {rf['rf_annual_max'] * 100:.2f}%"],
    ])

    if log.get("notes"):
        d.add_heading("8. Ghi chú tự động", 2)
        for n in log["notes"]:
            bullet(n)

    d.add_heading("9. Ghi chú thủ công của TV2", 2)
    if notes_path.exists():
        for title, lines in _sections(notes_path.read_text(encoding="utf-8")):
            for r in _md_tables_and_bullets(lines):
                bullet(_plain(" | ".join(r)))
    else:
        bullet("(chưa có – viết vào 05_Phu_luc/ghi_chu_thu_cong.md)")
    d.save(out)
    return out


def run() -> list[Path]:
    return [nhat_ky_docx(), tu_dien_xlsx()]


if __name__ == "__main__":
    for p in run():
        print("Đã xuất", p)
